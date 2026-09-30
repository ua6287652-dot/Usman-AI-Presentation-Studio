import io
import json
import os
import re
from typing import Any, Dict, List, Optional
from urllib.parse import quote

import requests

import streamlit as st
from groq import Groq
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

st.set_page_config(page_title="Usman AI Presentation Studio", page_icon="🎤", layout="wide")
MODEL = "openai/gpt-oss-20b"

THEMES = {
    "Midnight Blue": {"bg":"0B1220","panel":"111C33","card":"17243F","accent":"6D5DFB","accent2":"38BDF8","text":"F8FAFC","muted":"B8C2D9"},
    "Professional": {"bg":"F4F7FB","panel":"FFFFFF","card":"EEF3F9","accent":"2563EB","accent2":"0EA5E9","text":"0F172A","muted":"64748B"},
    "Emerald": {"bg":"071A16","panel":"0D2922","card":"12372E","accent":"10B981","accent2":"34D399","text":"F0FDF4","muted":"A7C7BB"},
    "Warm Sunset": {"bg":"24120D","panel":"351A12","card":"492319","accent":"F97316","accent2":"F43F5E","text":"FFF7ED","muted":"D7B6A7"},
    "Minimal Gray": {"bg":"171717","panel":"242424","card":"303030","accent":"A3A3A3","accent2":"D4D4D4","text":"FAFAFA","muted":"B3B3B3"},
}
PRESENTATION_TYPES = ["Educational","Business / Corporate","Pitch Deck","Project Proposal","Research","Workshop / Training","Marketing","Portfolio","General"]
AUDIENCES = ["Students","Teachers / Professors","Business Professionals","Clients / Customers","Investors","General Audience","Technical Audience","Management / Executives"]
LANGUAGES = ["English","Urdu","Roman Urdu","Arabic","Spanish","French","German"]
SLIDE_COUNTS = [5, 6, 7, 8, 10, 12, 15]
W, H = 13.333, 7.5


def rgb(v):
    v=v.replace("#","")
    return RGBColor(int(v[:2],16), int(v[2:4],16), int(v[4:6],16))


def api_key():
    try:
        key = st.secrets.get("GROQ_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GROQ_API_KEY", "")


def image_api_key():
    try:
        key = st.secrets.get("POLLINATIONS_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("POLLINATIONS_API_KEY", "")


def schema(n):
    return {"type":"object","additionalProperties":False,"properties":{
        "presentation_title":{"type":"string"},"subtitle":{"type":"string"},"overview":{"type":"string"},
        "slides":{"type":"array","minItems":n,"maxItems":n,"items":{"type":"object","additionalProperties":False,
            "properties":{"title":{"type":"string"},"subtitle":{"type":"string"},"bullets":{"type":"array","minItems":2,"maxItems":6,"items":{"type":"string"}},"speaker_notes":{"type":"string"}},
            "required":["title","subtitle","bullets","speaker_notes"]}},
    },"required":["presentation_title","subtitle","overview","slides"]}


def generate(topic, language, count, ptype, audience, theme, extra):
    content_count = max(1, count - 1)
    key=api_key()
    if not key: raise RuntimeError("GROQ_API_KEY is missing. Add it in Streamlit Cloud → Settings → Secrets.")
    client=Groq(api_key=key)
    prompt=f"""Create a professional presentation. Topic: {topic}. Language: {language}. Total PowerPoint slides including the title slide: {count}. The AI content slides to generate are exactly {content_count}. Type: {ptype}. Audience: {audience}. Visual theme: {theme}. Extra instructions: {extra or 'None'}. Return exactly {content_count} content slides. The application will use the first slide as the title slide, so the final PowerPoint must contain exactly {count} slides. Make the flow logical from introduction/context through main content to conclusion. Every slide needs a concise title, optional subtitle, 2-6 useful bullets, and speaker notes. Keep bullets concise enough for PowerPoint. Adapt depth to the audience. Use the requested language. Do not invent statistics, citations, or sources. Do not use markdown fences."""
    res=client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role":"system","content":"You are an expert presentation architect. Create accurate, clear, audience-appropriate presentation content."},
            {"role":"user","content":prompt},
        ],
        temperature=0.55,
        max_completion_tokens=7000,
        response_format={"type":"json_schema","json_schema":{"name":"presentation","strict":True,"schema":schema(content_count)}},
    )
    text=res.choices[0].message.content
    if not text: raise RuntimeError("Groq returned an empty response.")
    if text.strip().startswith("```"): text=re.sub(r"^```(?:json)?\s*|\s*```$","",text.strip(),flags=re.I)
    data=json.loads(text)
    if len(data.get("slides",[])) != content_count: raise RuntimeError(f"AI returned {len(data.get('slides',[]))} content slides instead of {content_count}.")
    return data


def make_image_prompt(title: str, subtitle: str, bullets: List[str], topic: str, theme_name: str) -> str:
    details = "; ".join(bullets[:4])
    return (
        f"Create a professional 16:9 presentation illustration about '{title}'. "
        f"Overall topic: {topic}. Context: {subtitle}. Key ideas: {details}. "
        f"Visual style: modern {theme_name} corporate editorial illustration, clean composition, "
        "high quality, realistic or polished 3D visual, strong depth, professional lighting, "
        "visually understandable for a presentation slide. No text, no letters, no words, "
        "no logos, no watermark, no charts with labels."
    )


def generate_topic_image(prompt: str, width: int = 1024, height: int = 576) -> bytes:
    key = image_api_key()
    if not key:
        raise RuntimeError(
            "POLLINATIONS_API_KEY is missing. Add it in Streamlit Cloud → Settings → Secrets "
            "to enable AI-generated slide images."
        )
    encoded = quote(prompt, safe="")
    url = (
        f"https://gen.pollinations.ai/image/{encoded}"
        f"?model=flux&width={width}&height={height}&nologo=true&private=true&safe=true"
    )
    response = requests.get(
        url,
        headers={"Authorization": f"Bearer {key}"},
        timeout=120,
    )
    response.raise_for_status()
    content_type = response.headers.get("content-type", "")
    if "image" not in content_type and not response.content.startswith(b"\xff\xd8") and not response.content.startswith(b"\x89PNG"):
        raise RuntimeError("The image service did not return a valid image.")
    return response.content


def generate_all_images(data: Dict[str, Any], topic: str, theme_name: str) -> List[Optional[bytes]]:
    images: List[Optional[bytes]] = []
    total = len(data.get("slides", [])) + 1
    progress = st.progress(0, text="Generating topic-wise AI images...")
    errors = []

    title_prompt = make_image_prompt(
        data.get("presentation_title", topic),
        data.get("subtitle", ""),
        [data.get("overview", "")],
        topic,
        theme_name,
    )
    prompts = [title_prompt]
    for slide in data.get("slides", []):
        prompts.append(
            make_image_prompt(
                slide.get("title", ""),
                slide.get("subtitle", ""),
                slide.get("bullets", []),
                topic,
                theme_name,
            )
        )

    for i, prompt in enumerate(prompts, start=1):
        try:
            images.append(generate_topic_image(prompt))
        except Exception as exc:
            images.append(None)
            errors.append(f"Slide {i}: {exc}")
        progress.progress(i / total, text=f"Generating AI image {i} of {total}...")

    progress.empty()
    if errors:
        st.warning("Some slide images could not be generated. The PowerPoint was still created. " + " | ".join(errors[:3]))
    return images


def textbox(slide,l,t,w,h,text,size=20,bold=False,color="FFFFFF",align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP):
    sh=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text or ""); r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return sh


def bg(slide,th):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=rgb(th["bg"])
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(W),Inches(.14)); bar.fill.solid(); bar.fill.fore_color.rgb=rgb(th["accent"]); bar.line.fill.background()


def notes(slide,text):
    if text:
        try: slide.notes_slide.notes_text_frame.text=text
        except Exception: pass


def add_image_to_slide(slide, image_bytes: Optional[bytes], left: float, top: float, width: float, height: float, th: Dict[str, str]):
    if image_bytes:
        slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(left), Inches(top), width=Inches(width), height=Inches(height))
    else:
        fallback = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        fallback.fill.solid()
        fallback.fill.fore_color.rgb = rgb(th["panel"])
        fallback.line.color.rgb = rgb(th["accent"])
        textbox(
            slide, left + 0.25, top + height / 2 - 0.25, width - 0.5, 0.5,
            "AI image unavailable", 12, True, th["muted"], PP_ALIGN.CENTER
        )


def make_ppt(data, theme_name, images: Optional[List[Optional[bytes]]] = None):
    th = THEMES[theme_name]
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    images = images or []

    # Title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, th)
    add_image_to_slide(s, images[0] if len(images) > 0 else None, 8.0, 0.85, 4.65, 5.75, th)
    textbox(s, .75, 1.0, 6.7, 1.35, data["presentation_title"], 32, True, th["text"])
    textbox(s, .78, 2.55, 6.45, .75, data["subtitle"], 18, False, th["accent2"])
    textbox(s, .78, 3.55, 6.55, 1.55, data["overview"], 15, False, th["muted"])
    textbox(s, .78, 6.35, 6.8, .3, "Created with Usman AI Presentation Studio", 11, True, th["accent"])
    notes(s, "Opening slide. Introduce the topic and explain what the audience will learn.")

    total = len(data["slides"]) + 1
    for num, d in enumerate(data["slides"], 2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg(s, th)
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.62), Inches(1.0), Inches(.08), Inches(4.9))
        line.fill.solid(); line.fill.fore_color.rgb = rgb(th["accent"]); line.line.fill.background()

        textbox(s, .95, .68, 7.1, .65, d["title"], 26, True, th["text"])
        if d.get("subtitle"):
            textbox(s, .97, 1.34, 7.0, .42, d["subtitle"], 11, False, th["accent2"])

        # Content panel
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.95), Inches(1.95), Inches(7.05), Inches(4.65))
        card.fill.solid(); card.fill.fore_color.rgb = rgb(th["card"]); card.line.fill.background()
        y = 2.28
        for bullet in d.get("bullets", [])[:6]:
            dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.25), Inches(y+.06), Inches(.13), Inches(.13))
            dot.fill.solid(); dot.fill.fore_color.rgb = rgb(th["accent"]); dot.line.fill.background()
            textbox(s, 1.55, y, 6.05, .62, bullet, 15, False, th["text"])
            y += .70

        # Topic-wise AI image
        image_index = num - 1
        add_image_to_slide(s, images[image_index] if image_index < len(images) else None, 8.35, 1.95, 4.15, 4.65, th)
        textbox(s, .65, 7.05, 10.7, .22, "Usman AI Presentation Studio", 8, False, th["muted"])
        textbox(s, 11.65, 7.02, 1, .25, f"{num}/{total}", 9, False, th["muted"], PP_ALIGN.RIGHT)
        notes(s, d.get("speaker_notes", ""))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def css(theme):
    th=THEMES[theme]
    st.markdown(f"""<style>
    .stApp{{background:linear-gradient(135deg,#{th['bg']} 0%,#0f172a 100%);}}
    [data-testid='stSidebar']{{background:#{th['panel']};}}
    .brand{{padding:15px 5px 8px}} .brand-title{{font-size:25px;font-weight:800;color:#{th['text']}}} .brand-sub{{color:#{th['muted']};font-size:12px}}
    .hero{{padding:28px;border-radius:22px;background:linear-gradient(135deg,#{th['accent']},#{th['accent2']});color:white;margin-bottom:18px;box-shadow:0 16px 45px rgba(0,0,0,.18)}}
    .hero h1{{margin:0;font-size:36px;line-height:1.1}} .hero p{{margin:10px 0 0;opacity:.92}}
    .title{{color:#{th['text']};font-size:22px;font-weight:750;margin:8px 0 10px}}
    .mini{{padding:18px;border-radius:18px;background:#{th['panel']};border:1px solid rgba(255,255,255,.08);min-height:118px}}
    .mini h4{{color:#{th['text']};margin:0 0 6px}} .mini p{{color:#{th['muted']};margin:0;font-size:13px}}
    div.stButton>button{{border-radius:12px;font-weight:700}}
    footer{{visibility:hidden}}
    </style>""",unsafe_allow_html=True)


if "data" not in st.session_state: st.session_state.data=None
if "ppt" not in st.session_state: st.session_state.ppt=None
if "images" not in st.session_state: st.session_state.images=[]

with st.sidebar:
    st.markdown('<div class="brand"><div class="brand-title">🎤 Usman AI</div><div class="brand-sub">Presentation Studio</div></div>',unsafe_allow_html=True)
    st.markdown("---")
    theme=st.selectbox("🎨 Visual Theme",list(THEMES.keys()))
    st.markdown("---")
    st.markdown("### 🖼️ AI Slide Images")
    use_images = st.checkbox("Generate topic-wise AI image for every slide", value=True)
    if use_images:
        st.caption("Uses Pollinations AI image generation. Add POLLINATIONS_API_KEY to Streamlit Secrets.")
    st.markdown("---")
    st.markdown("### 💡 Workflow")
    st.markdown("**1. Define** → **2. Generate** → **3. Review** → **4. Download**")
    st.caption("Groq + Streamlit + python-pptx")
css(theme)

st.markdown('<div class="hero"><h1>Usman AI Presentation Studio ✨</h1><p>Create structured, professional and editable PowerPoint presentations from your requirements.</p></div>',unsafe_allow_html=True)
cols=st.columns(4)
for col,icon,title,desc in zip(cols,["🤖","🎯","🎨","📁"],["AI Powered","Customizable","Visual Themes","Editable PPTX"],["Generate structured slide content.","Control language, audience and type.","Choose a presentation style.","Download a real editable PowerPoint."]):
    with col: st.markdown(f'<div class="mini"><div style="font-size:24px">{icon}</div><h4>{title}</h4><p>{desc}</p></div>',unsafe_allow_html=True)

st.markdown('<div class="title">Build Your Presentation</div>',unsafe_allow_html=True)
left,right=st.columns([1.05,.95],gap="large")
with left:
    topic=st.text_area("Presentation Topic *",placeholder="Example: Artificial Intelligence in Education",height=100,max_chars=500)
    a,b=st.columns(2)
    with a:
        language=st.selectbox("Language",LANGUAGES)
        count=st.selectbox("Number of Slides",SLIDE_COUNTS,index=3)
        ptype=st.selectbox("Presentation Type",PRESENTATION_TYPES)
    with b:
        audience=st.selectbox("Target Audience",AUDIENCES)
        visual=st.selectbox("Visual Theme",list(THEMES.keys()),index=list(THEMES.keys()).index(theme))
        extra=st.text_area("Additional Instructions",placeholder="Example: Keep it simple and include practical examples.",height=100)
    go=st.button("✨ Generate Presentation",type="primary",use_container_width=True)
with right:
    st.markdown('<div class="title">Presentation Preview</div>',unsafe_allow_html=True)
    if st.session_state.data:
        d=st.session_state.data; st.success(f"Generated {len(d['slides'])+1} slides")
        st.markdown(f"### {d['presentation_title']}"); st.caption(d['subtitle']); st.write(d['overview'])
        with st.expander("📑 View Slide Outline",expanded=True):
            for i,s in enumerate(d["slides"],1):
                st.markdown(f"**Slide {i}: {s['title']}**")
                for bullet in s["bullets"]: st.markdown(f"- {bullet}")
        if st.session_state.images:
            with st.expander("🖼️ Preview Topic-wise AI Images", expanded=False):
                for idx, image in enumerate(st.session_state.images, start=1):
                    if image:
                        st.image(image, caption=f"Slide {idx} image", use_container_width=True)

        if st.session_state.ppt:
            st.download_button("📥 Download Editable PowerPoint (.pptx)",st.session_state.ppt,"usman_ai_presentation.pptx","application/vnd.openxmlformats-officedocument.presentationml.presentation",use_container_width=True)
    else: st.info("Your generated presentation will appear here. Fill in the requirements and click Generate Presentation.")

if go:
    if not topic.strip(): st.warning("Please enter a presentation topic.")
    else:
        with st.spinner("🤖 Creating your presentation..."):
            try:
                d=generate(topic.strip(),language,count,ptype,audience,visual,extra.strip())
                images = generate_all_images(d, topic.strip(), visual) if use_images else []
                st.session_state.data=d
                st.session_state.images=images
                st.session_state.ppt=make_ppt(d,visual,images)
                st.success("🎉 Presentation and topic-wise images generated successfully!")
                st.rerun()
            except Exception as e: st.error(f"Could not generate presentation: {e}")

st.markdown("---")
st.markdown('<div class="title">🚀 Project Workflow</div>',unsafe_allow_html=True)
x,y,z=st.columns(3)
for col,title,desc in [(x,"Generative AI Workflow","Requirements → AI planning → structured slide content → PowerPoint."),(y,"Real Output","The result is an editable .pptx file, not just plain text."),(z,"Vibe Coding Friendly","Deploy directly from GitHub to Streamlit Cloud.")]:
    with col: st.markdown(f'<div class="mini"><h4>{title}</h4><p>{desc}</p></div>',unsafe_allow_html=True)
